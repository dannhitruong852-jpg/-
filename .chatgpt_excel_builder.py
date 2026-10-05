# -*- coding: utf-8 -*-
import csv, io, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(ROOT, ".chatgpt_excel_source.json"), "r", encoding="utf-8") as f:
    payload = json.load(f)

out_name = "2021-2026浙大MAP单选考点定位.xlsx"
out_path = os.path.join(ROOT, out_name)

wb = Workbook()
wb.remove(wb.active)

title_fill = PatternFill("solid", fgColor="1F4E78")
header_fill = PatternFill("solid", fgColor="D9EAF7")
note_fill = PatternFill("solid", fgColor="F2F2F2")
thin_gray = Side(style="thin", color="D9D9D9")

for sheet_obj in payload["sheets"]:
    ws = wb.create_sheet(sheet_obj["name"][:31])
    rows = list(csv.reader(io.StringIO(sheet_obj["csv"])))
    # Parsed text has a synthetic index column and a synthetic header row.
    data_rows = rows[1:]
    for r_idx, row in enumerate(data_rows, start=1):
        vals = row[1:5]
        while len(vals) < 4:
            vals.append("")
        for c_idx, val in enumerate(vals, start=1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Microsoft YaHei", size=10.5)

    # Title / guidance rows
    for rr in (1, 2, 3):
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=4)
    ws["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A2"].font = Font(name="Microsoft YaHei", size=10.5, italic=True, color="404040")
    ws["A3"].font = Font(name="Microsoft YaHei", size=10.5, color="666666")
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 32

    # Header row is row 5 in the original layout.
    for cell in ws[5]:
        cell.font = Font(name="Microsoft YaHei", size=10.5, bold=True, color="1F1F1F")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)
    ws.row_dimensions[5].height = 26

    # Notes at bottom
    max_row = ws.max_row
    for rr in range(max(1, max_row - 1), max_row + 1):
        for cell in ws[rr]:
            cell.fill = note_fill
            cell.font = Font(name="Microsoft YaHei", size=9.5, color="666666")
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [18, 42, 72, 78]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A5:D{max_row}"

wb.save(out_path)
print(out_path)
